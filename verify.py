from pptx import Presentation

prs = Presentation("D:/线束业务培训/output/青岛弘博电缆_新业务入职培训.pptx")
print("Total slides:", len(prs.slides))
slide_count = 0
for slide in prs.slides:
    slide_count += 1
    if slide_count > 10:
        break
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            text = shape.text_frame.text.replace("\n", " ").strip()[:60]
            if text:
                texts.append(text)
    first_text = texts[0] if texts else "(no text)"
    print(f"Slide {slide_count}: {first_text}")
